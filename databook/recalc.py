"""Recalculate every formula in an .xlsx with LibreOffice headless, then scan for errors.

openpyxl writes formulas without cached values. LibreOffice opens the workbook, runs
calculateAll() through a Basic macro installed in a throwaway user profile, and saves it
in place, so every formula cell then carries its computed value.

Usage: python databook/recalc.py path/to/workbook.xlsx
Prints JSON: status, total_formulas, total_errors, error cells. Exit code 1 on failure.
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""

ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")


def _env() -> dict:
    env = os.environ.copy()
    env["SAL_USE_VCLPLUGIN"] = "svp"
    return env


def recalc(path: str, timeout: int = 120) -> dict:
    target = Path(path).resolve()
    if not shutil.which("soffice"):
        return {"error": "soffice not found; LibreOffice is required"}
    with tempfile.TemporaryDirectory(prefix="lo-profile-") as profile:
        url = Path(profile).as_uri()
        subprocess.run(["soffice", "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                       env=_env(), capture_output=True, timeout=timeout)
        macro_dir = Path(profile) / "user" / "basic" / "Standard"
        if not macro_dir.exists():
            return {"error": "LibreOffice did not create a user profile"}
        (macro_dir / "Module1.xba").write_text(MACRO, encoding="utf-8")
        before = target.stat().st_mtime_ns
        result = subprocess.run(
            ["soffice", "--headless", "--norestore", f"-env:UserInstallation={url}",
             "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application", str(target)],
            env=_env(), capture_output=True, text=True, timeout=timeout)
        if result.returncode != 0:
            return {"error": f"soffice exited {result.returncode}: {result.stderr.strip()}"}
        if target.stat().st_mtime_ns == before:
            return {"error": "LibreOffice did not rewrite the file; nothing was recalculated"}
    return scan(str(target))


def scan(path: str) -> dict:
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    errors, n_formulas = [], 0
    for ws in formulas.worksheets:
        vs = values[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_formulas += 1
                    v = vs[c.coordinate].value
                    if v is None:
                        errors.append(f"{ws.title}!{c.coordinate} (no cached value)")
                    elif isinstance(v, str) and any(e in v for e in ERRORS):
                        errors.append(f"{ws.title}!{c.coordinate} {v}")
    return {"status": "success" if not errors else "errors_found", "total_formulas": n_formulas,
            "total_errors": len(errors), "errors": errors[:100]}


if __name__ == "__main__":
    out = recalc(sys.argv[1])
    print(json.dumps(out, indent=2))
    sys.exit(1 if "error" in out or out.get("status") != "success" else 0)
