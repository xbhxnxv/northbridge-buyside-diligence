"""Tie-outs and integrity checks for the Northbridge pipeline.

Run after `python run_all.py` (or `python run_all.py --only sql`), which builds
data/northbridge.duckdb. Every test reads the database read-only.
"""

from __future__ import annotations

import csv
import json
from decimal import Decimal
from pathlib import Path

import duckdb
import pytest

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
DB_PATH = ROOT / "data" / "northbridge.duckdb"

TABLES = ["customers", "products", "subscriptions", "invoices", "costs", "management_accounts"]


@pytest.fixture(scope="session")
def con():
    if not DB_PATH.exists():
        pytest.skip("data/northbridge.duckdb not built; run `python run_all.py` first")
    c = duckdb.connect(str(DB_PATH), read_only=True)
    yield c
    c.close()


def scalar(con, sql: str):
    return con.execute(sql).fetchone()[0]


def csv_rows(name: str) -> list[dict]:
    with open(RAW / f"{name}.csv", newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# ---------------------------------------------------------------------------
# Step 0: the data room is the one the brief specifies
# ---------------------------------------------------------------------------

EXPECTED_ROWS = {
    "customers": 2_192,
    "products": 6,
    "subscriptions": 3_502,
    "invoices": 92_732,
    "costs": 192,
    "management_accounts": 48,
}

# Published in the project brief, rounded to the nearest pound.
EXPECTED_MGMT_REVENUE = {2022: 28_232_152, 2023: 32_174_370, 2024: 37_261_399, 2025: 45_771_419}


@pytest.mark.parametrize("name", TABLES)
def test_data_room_row_counts(name):
    assert len(csv_rows(name)) == EXPECTED_ROWS[name]


def test_data_room_management_revenue():
    totals: dict[int, Decimal] = {}
    for r in csv_rows("management_accounts"):
        y = int(r["month"][:4])
        totals[y] = totals.get(y, Decimal(0)) + Decimal(r["total_revenue"])
    assert {y: round(v) for y, v in totals.items()} == EXPECTED_MGMT_REVENUE


# ---------------------------------------------------------------------------
# Step 2.1: the typed load drops nothing and rounds nothing
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("name", TABLES)
def test_load_row_counts_match_csv(con, name):
    assert scalar(con, f"SELECT count(*) FROM raw_{name}") == len(csv_rows(name))


MONEY_COLUMNS = {
    "invoices": ["amount"],
    "subscriptions": ["monthly_price"],
    "costs": ["hosting_cost", "support_staff_cost", "third_party_licence_cost", "total_cost_of_delivery"],
    "management_accounts": ["revenue_core_platform", "revenue_analytics_add-on", "revenue_payments_module",
                            "revenue_implementation_services", "total_revenue", "cost_of_sales", "gross_profit"],
}


@pytest.mark.parametrize("name,column", [(t, c) for t, cols in MONEY_COLUMNS.items() for c in cols])
def test_load_money_totals_match_csv_text(con, name, column):
    """Sum of the CSV text parsed as exact decimals equals the sum of the loaded DECIMAL column."""
    expected = sum(Decimal(r[column]) for r in csv_rows(name))
    loaded_column = "revenue_analytics_addon" if column == "revenue_analytics_add-on" else column
    assert scalar(con, f"SELECT sum({loaded_column}) FROM raw_{name}") == expected


def test_load_no_null_dates_introduced(con):
    """Blank end_date is the only permitted null date: every other date parsed."""
    assert scalar(con, "SELECT count(*) FROM raw_invoices WHERE invoice_date IS NULL") == 0
    assert scalar(con, "SELECT count(*) FROM raw_customers WHERE signup_date IS NULL") == 0
    assert scalar(con, "SELECT count(*) FROM raw_subscriptions WHERE start_date IS NULL") == 0
    blank_end = sum(1 for r in csv_rows("subscriptions") if r["end_date"] == "")
    assert scalar(con, "SELECT count(*) FROM raw_subscriptions WHERE end_date IS NULL") == blank_end


# ---------------------------------------------------------------------------
# Step 2.2: profile integrity
# ---------------------------------------------------------------------------

def test_profile_primary_keys_unique(con):
    assert scalar(con, "SELECT count(*) FROM prof_primary_keys WHERE row_count <> distinct_keys OR null_keys > 0") == 0


def test_profile_issue_register_complete(con):
    ids = [r[0] for r in con.execute("SELECT issue_id FROM prof_issues ORDER BY issue_id").fetchall()]
    assert ids == [f"P{i:02d}" for i in range(1, 12)]


def test_profile_issue_register_years_add_up(con):
    bad = scalar(con, """
        SELECT count(*) FROM prof_issues
        WHERE abs(amount_net - (net_2022 + net_2023 + net_2024 + net_2025)) > 0.005
    """)
    assert bad == 0


# ---------------------------------------------------------------------------
# Step 2.3: cleaning
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("name", TABLES)
def test_clean_rows_reconcile_to_raw(con, name):
    """raw rows = clean rows + removed rows, per table."""
    raw = scalar(con, f"SELECT count(*) FROM raw_{name}")
    clean = scalar(con, f"SELECT count(*) FROM clean_{name}")
    removed = scalar(con, "SELECT count(*) FROM clean_removed_invoices") if name == "invoices" else 0
    assert raw == clean + removed


def test_removed_rows_are_exactly_the_duplicate_copies(con):
    assert scalar(con, "SELECT count(*) FROM clean_removed_invoices") == 40
    assert scalar(con, """
        SELECT count(*) FROM clean_removed_invoices r
        WHERE NOT EXISTS (SELECT 1 FROM clean_invoices k
                          WHERE k.invoice_id = r.kept_invoice_id AND k.customer_id = r.customer_id
                            AND k.product_id = r.product_id AND k.invoice_date = r.invoice_date
                            AND k.amount = r.amount AND k.status = r.status)
    """) == 0
    # no exact duplicates remain
    assert scalar(con, """
        SELECT count(*) FROM (SELECT 1 FROM clean_invoices
                              GROUP BY customer_id, product_id, invoice_date, amount, revenue_type, status
                              HAVING count(*) > 1)
    """) == 0


def test_p04_amount_removed(con):
    assert scalar(con, "SELECT sum(amount) FROM clean_removed_invoices") == Decimal("96957.26")


def test_every_credit_note_linked_to_its_original(con):
    assert scalar(con, "SELECT count(*) FROM clean_invoices WHERE is_credit_note AND credit_note_original_id IS NULL") == 0
    bad = scalar(con, """
        SELECT count(*) FROM clean_invoices cn
        LEFT JOIN clean_invoices o ON o.invoice_id = cn.credit_note_original_id
        WHERE cn.is_credit_note
          AND (o.invoice_id IS NULL OR o.is_credit_note OR o.customer_id <> cn.customer_id
               OR o.product_id <> cn.product_id OR o.invoice_date <> cn.invoice_date OR o.amount <> -cn.amount
               OR o.credited_by_id <> cn.invoice_id)
    """)
    assert bad == 0
    # one-to-one: no original is linked by two credit notes
    assert scalar(con, """SELECT count(*) FROM (SELECT credit_note_original_id FROM clean_invoices
                          WHERE is_credit_note GROUP BY 1 HAVING count(*) > 1)""") == 0


def test_cleaning_actions_explain_the_revenue_change(con):
    """raw invoice total + revenue impact of all cleaning actions = cube net (current setting)."""
    raw = scalar(con, "SELECT sum(amount) FROM raw_invoices")
    impact = scalar(con, "SELECT sum(revenue_impact) FROM cleaning_actions")
    cube = scalar(con, "SELECT sum(net_revenue) FROM fact_revenue_monthly")
    assert raw + impact == cube


def test_customers_never_dropped_and_industry_filled(con):
    assert scalar(con, "SELECT count(*) FROM clean_customers WHERE industry IS NULL") == 0
    assert scalar(con, "SELECT count(*) FROM clean_customers WHERE industry = 'Unknown'") == \
        scalar(con, "SELECT count(*) FROM raw_customers WHERE industry IS NULL")


# ---------------------------------------------------------------------------
# Step 2.4: revenue cube
# ---------------------------------------------------------------------------

def test_cube_net_equals_clean_invoice_net_total_and_by_year(con):
    rows = con.execute("""
        WITH inv AS (
            SELECT year(month_start) AS y,
                   sum(invoice_amount((SELECT anomaly_treatment FROM cfg_settings),
                                      amount_as_reported, amount_flipped, amount_excluded)) AS net
            FROM clean_invoices GROUP BY 1),
        cube AS (SELECT year(month_start) AS y, sum(net_revenue) AS net FROM fact_revenue_monthly GROUP BY 1)
        SELECT inv.y, inv.net, cube.net FROM inv FULL JOIN cube USING (y) ORDER BY 1
    """).fetchall()
    assert [r[0] for r in rows] == [2022, 2023, 2024, 2025]
    for y, inv_net, cube_net in rows:
        assert inv_net == cube_net, y
    assert sum(r[1] for r in rows) == scalar(con, "SELECT sum(net_revenue) FROM fact_revenue_monthly")


def test_cube_row_identities(con):
    assert scalar(con, """
        SELECT count(*) FROM fact_revenue_monthly
        WHERE recurring_net + oneoff_net <> net_revenue
           OR gross_billings + credit_notes <> net_revenue
           OR oneoff_signup_net + oneoff_other_net <> oneoff_net
    """) == 0


@pytest.mark.parametrize("table,grain", [
    ("fact_revenue_monthly", "customer_id, product_id, month_start"),
    ("fact_mrr_monthly", "customer_id, product_id, month_start"),
    ("dim_customer", "customer_id"),
    ("dim_product", "product_id"),
    ("dim_date", "month_start"),
])
def test_no_duplicate_grain_rows(con, table, grain):
    assert scalar(con, f"SELECT count(*) FROM (SELECT {grain} FROM {table} GROUP BY {grain} HAVING count(*) > 1)") == 0


@pytest.mark.parametrize("fact", ["fact_revenue_monthly", "fact_mrr_monthly"])
def test_cube_keys_exist_in_dimensions(con, fact):
    assert scalar(con, f"SELECT count(*) FROM {fact} f WHERE NOT EXISTS (SELECT 1 FROM dim_customer d WHERE d.customer_id = f.customer_id)") == 0
    assert scalar(con, f"SELECT count(*) FROM {fact} f WHERE NOT EXISTS (SELECT 1 FROM dim_product d WHERE d.product_id = f.product_id)") == 0
    assert scalar(con, f"SELECT count(*) FROM {fact} f WHERE NOT EXISTS (SELECT 1 FROM dim_date d WHERE d.month_start = f.month_start)") == 0


def test_mrr_is_recurring_billings_before_credit_notes(con):
    mrr = scalar(con, "SELECT sum(mrr) FROM fact_mrr_monthly")
    billings = scalar(con, "SELECT sum(gross_billings) FROM fact_revenue_monthly WHERE revenue_type = 'recurring'")
    assert mrr == billings


def test_anomaly_switch_moves_revenue_by_exactly_the_anomaly_amounts(con):
    rows = con.execute("""
        WITH a AS (SELECT year(month_start) AS y, sum(net_revenue) AS net FROM revenue_cube('as_reported') GROUP BY 1),
             f AS (SELECT year(month_start) AS y, sum(net_revenue) AS net FROM revenue_cube('flipped') GROUP BY 1),
             e AS (SELECT year(month_start) AS y, sum(net_revenue) AS net FROM revenue_cube('excluded') GROUP BY 1),
             n AS (SELECT year(month_start) AS y, sum(amount) AS anomaly FROM clean_invoices WHERE anomaly_flag GROUP BY 1)
        SELECT a.y, a.net, f.net, e.net, coalesce(n.anomaly, 0)
        FROM a JOIN f USING (y) JOIN e USING (y) LEFT JOIN n USING (y) ORDER BY 1
    """).fetchall()
    assert len(rows) == 4
    for y, as_rep, flipped, excluded, anomaly in rows:
        assert flipped - as_rep == -2 * anomaly, y     # anomaly amounts are negative
        assert excluded - as_rep == -anomaly, y
    assert scalar(con, "SELECT count(*) FROM clean_invoices WHERE anomaly_flag AND amount >= 0") == 0


def test_exports_match_tables(con):
    clean_dir = ROOT / "data" / "clean"
    for t in ["fact_revenue_monthly", "fact_mrr_monthly", "dim_customer", "dim_product", "dim_date"]:
        with open(clean_dir / f"{t}.csv", newline="", encoding="utf-8") as f:
            n = sum(1 for _ in f) - 1
        assert n == scalar(con, f"SELECT count(*) FROM {t}"), t


def test_source_to_cube_walk_adds_up():
    path = ROOT / "outputs" / "tables" / "step2_source_to_cube_walk.csv"
    if not path.exists():
        pytest.skip("walk not built; run the notebooks")
    with open(path, newline="") as f:
        for r in csv.DictReader(f):
            total = sum(Decimal(r[c]) for c in ["raw_gross_billings", "less_duplicates", "anomaly_treatment", "credit_notes"])
            assert abs(total - Decimal(r["cube_net"])) < Decimal("0.005"), r["year"]


def test_qa_log_ids_sequential(con):
    ids = [r[0] for r in con.execute("SELECT qa_id FROM qa_log ORDER BY qa_id").fetchall()]
    assert ids == [f"Q{i:02d}" for i in range(1, len(ids) + 1)]


# ---------------------------------------------------------------------------
# Step 3: reconciliation (checks that hold whatever anomaly treatment is chosen)
# ---------------------------------------------------------------------------

def test_recon_walk_adds_up(con):
    bad = scalar(con, """
        SELECT count(*) FROM recon_walk
        WHERE raw_gross_billings + less_duplicates + anomaly_treatment + credit_notes <> cube_net
           OR cube_net + unreconciled_difference <> mgmt_total
    """)
    assert bad == 0
    assert scalar(con, "SELECT count(*) FROM recon_walk") == 4


def test_recon_walk_matches_management_accounts(con):
    assert scalar(con, "SELECT sum(mgmt_total) FROM recon_walk") == \
        scalar(con, "SELECT sum(total_revenue) FROM clean_management_accounts")


def test_costs_tie_to_cost_of_sales_every_month(con):
    assert scalar(con, "SELECT count(*) FROM recon_costs_monthly WHERE difference <> 0") == 0
    assert scalar(con, "SELECT count(*) FROM recon_costs_monthly") == 48


def test_options_grid_complete(con):
    assert scalar(con, "SELECT count(*) FROM recon_options_monthly") == 3 * 48


def test_product_lines_sum_to_cube_total(con):
    bad = scalar(con, """
        SELECT count(*) FROM (
            SELECT p.month, sum(p.cube_net) AS lines, any_value(r.cube_net) AS total
            FROM recon_product_line_monthly p JOIN recon_monthly r USING (month) GROUP BY p.month)
        WHERE lines <> total
    """)
    assert bad == 0


def test_anomaly_setting_is_excluded_and_ties_to_lines(con):
    """Step 3 decision D01: excluded is the only treatment that ties to the product lines every month."""
    assert scalar(con, "SELECT anomaly_treatment FROM cfg_settings") == "excluded"
    tied = dict(con.execute("SELECT treatment, months_tied_to_lines FROM recon_options_summary").fetchall())
    assert tied["excluded"] == 48
    assert tied["as_reported"] < 48 and tied["flipped"] < 48


def test_no_negative_mrr_rows(con):
    assert scalar(con, "SELECT count(*) FROM fact_mrr_monthly WHERE mrr <= 0") == 0


def test_every_month_labelled(con):
    assert scalar(con, "SELECT count(*) FROM recon_monthly") == 48
    assert scalar(con, "SELECT count(*) FROM recon_monthly WHERE status NOT IN ('tied', 'explained', 'unexplained') OR status IS NULL") == 0
    # tied and explained months have no difference; unexplained months carry an explanation
    assert scalar(con, "SELECT count(*) FROM recon_monthly WHERE status <> 'unexplained' AND difference <> 0") == 0
    assert scalar(con, "SELECT count(*) FROM recon_monthly WHERE status <> 'tied' AND explanation IS NULL") == 0
    assert scalar(con, "SELECT count(*) FROM recon_product_line_monthly WHERE status IS NULL") == 0


def test_annual_difference_equals_documented_unexplained_items(con):
    rows = con.execute("""
        SELECT a.year, a.difference, coalesce(u.amount, 0)
        FROM recon_annual a
        LEFT JOIN (SELECT year, sum(amount) AS amount FROM recon_unexplained_items GROUP BY 1) u USING (year)
    """).fetchall()
    assert len(rows) == 4
    for year, diff, items in rows:
        assert diff == items, year


def test_unexplained_items_are_the_management_line_gaps(con):
    """Every unexplained month is one where the cube ties to the product lines but not the reported total."""
    assert scalar(con, """
        SELECT count(*) FROM recon_unexplained_items u
        JOIN clean_management_accounts m USING (month)
        WHERE u.amount <> m.lines_less_total OR NOT m.lines_total_mismatch_flag
    """) == 0
    assert scalar(con, "SELECT count(*) FROM recon_unexplained_items") == \
        scalar(con, "SELECT count(*) FROM clean_management_accounts WHERE lines_total_mismatch_flag")


def test_product_lines_tie_every_month_under_chosen_setting(con):
    assert scalar(con, "SELECT count(*) FROM recon_product_line_monthly WHERE difference <> 0") == 0
    assert scalar(con, "SELECT count(*) FROM recon_product_line_monthly") == 4 * 48


def test_annual_reconciliation_within_tolerance_apart_from_documented_items(con):
    """After removing the documented unexplained items, every year ties to the penny."""
    assert scalar(con, """
        SELECT count(*) FROM recon_annual a
        WHERE abs(a.difference - coalesce((SELECT sum(amount) FROM recon_unexplained_items u WHERE u.year = a.year), 0)) >= 0.01
    """) == 0


# ---------------------------------------------------------------------------
# Step 4: core analyses (outputs/tables checked against the database)
# ---------------------------------------------------------------------------

TABLES_DIR = ROOT / "outputs" / "tables"


def table(name: str) -> list[dict]:
    path = TABLES_DIR / f"{name}.csv"
    if not path.exists():
        pytest.skip(f"{path.name} not built; run the notebooks")
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def close(a, b, tol=0.01) -> bool:
    return abs(float(a) - float(b)) <= tol


def test_arr_bridge_ties_every_year(con):
    comps = ["new", "cross_sell", "upsell", "price_increase", "downgrade", "contraction", "churn"]
    rows = table("4e_arr_bridge")
    assert [int(r["year"]) for r in rows] == [2023, 2024, 2025]
    for r in rows:
        y = int(r["year"])
        assert close(float(r["opening_arr"]) + sum(float(r[c]) for c in comps), r["closing_arr"]), y
        # opening and closing equal December MRR x 12 straight from the cube
        for col, yy in (("opening_arr", y - 1), ("closing_arr", y)):
            db = scalar(con, f"SELECT 12 * sum(mrr) FROM fact_mrr_monthly WHERE month_start = DATE '{yy}-12-01'")
            assert close(r[col], db), (y, col)
        # new and churn customer counts recomputed independently
        new, churned = con.execute(f"""
            WITH o AS (SELECT DISTINCT customer_id FROM fact_mrr_monthly WHERE month_start = DATE '{y - 1}-12-01'),
                 c AS (SELECT DISTINCT customer_id FROM fact_mrr_monthly WHERE month_start = DATE '{y}-12-01')
            SELECT (SELECT count(*) FROM c WHERE customer_id NOT IN (SELECT customer_id FROM o)),
                   (SELECT count(*) FROM o WHERE customer_id NOT IN (SELECT customer_id FROM c))
        """).fetchone()
        assert int(float(r["new_customers"])) == new and int(float(r["churned_customers"])) == churned, y


def test_nrr_grr_recomputed_independently(con):
    reported = {(int(r["year"]), r["basis"]): r for r in table("4d_nrr_grr")}
    for y in (2023, 2024, 2025):
        nrr, grr = con.execute(f"""
            WITH m AS (SELECT customer_id, month_start, sum(mrr) AS mrr FROM fact_mrr_monthly
                       WHERE month_start IN (DATE '{y - 1}-12-01', DATE '{y}-12-01') GROUP BY ALL),
                 o AS (SELECT customer_id, mrr FROM m WHERE month_start = DATE '{y - 1}-12-01'),
                 j AS (SELECT o.customer_id, o.mrr AS open_mrr, coalesce(c.mrr, 0) AS close_mrr
                       FROM o LEFT JOIN m c ON c.customer_id = o.customer_id AND c.month_start = DATE '{y}-12-01')
            SELECT 100 * sum(close_mrr) / sum(open_mrr), 100 * sum(least(close_mrr, open_mrr)) / sum(open_mrr) FROM j
        """).fetchone()
        r = reported[(y, "reported")]
        assert close(r["nrr_pct"], nrr, 1e-6) and close(r["grr_pct"], grr, 1e-6), y
        assert float(r["grr_pct"]) <= float(r["nrr_pct"]) + 1e-9
    assert {y for y, _ in reported} == {2023, 2024, 2025}   # no 2022: there is no December 2021


def test_concentration_shares(con):
    rows = table("4b_top_n_shares")
    by_year: dict[int, list] = {}
    for r in rows:
        assert 0 < float(r["share_of_net_pct"]) <= 100 and 0 < float(r["share_of_recurring_pct"]) <= 100
        assert close(100 * float(r["net_revenue_top_n"]) / float(r["net_revenue_total"]), r["share_of_net_pct"], 1e-6)
        by_year.setdefault(int(r["year"]), []).append(r)
    for y, rs in by_year.items():
        rs.sort(key=lambda r: int(r["top_n"]))
        shares = [float(r["share_of_net_pct"]) for r in rs]
        assert shares == sorted(shares), y          # top 5 >= top 1, and so on
        total = scalar(con, f"SELECT sum(net_revenue) FROM fact_revenue_monthly WHERE year(month_start) = {y}")
        top1 = scalar(con, f"""SELECT max(s) FROM (SELECT sum(net_revenue) AS s FROM fact_revenue_monthly
                               WHERE year(month_start) = {y} GROUP BY customer_id)""")
        assert close(rs[0]["net_revenue_total"], total) and close(rs[0]["net_revenue_top_n"], top1), y


def test_revenue_bridge_equals_cube_change(con):
    parts = ["new_customers", "full_year_effect_of_prior_year_additions", "cross_sell", "upsell", "price_increase", "billing_gaps",
             "downgrade", "contraction", "churn", "credit_notes_change", "signup_implementation_change", "other_implementation_change"]
    for r in table("4e_revenue_bridge"):
        y = int(r["year"])
        cube = {yy: scalar(con, f"SELECT sum(net_revenue) FROM fact_revenue_monthly WHERE year(month_start) = {yy}") for yy in (y - 1, y)}
        assert close(sum(float(r[p]) for p in parts), cube[y] - cube[y - 1]), y


def test_revenue_mix_adds_up(con):
    for r in table("4a_revenue_mix"):
        assert close(float(r["recurring"]) + float(r["oneoff"]), r["net_revenue"])
        assert close(float(r["signup_implementation"]) + float(r["other_implementation"]), r["oneoff"])
        assert close(r["net_revenue"], scalar(con, f"SELECT sum(net_revenue) FROM fact_revenue_monthly WHERE year(month_start) = {r['year']}"))


def test_uplift_factors_quoted_in_docs_match_data():
    k = {int(r["year"]): float(r["k"]) for r in table("4a_uplift_factors")}
    svi = {r["year_end"][:4]: float(r["invoiced_over_subscription"]) for r in table("4a_sub_vs_invoice_mrr")}
    text = (ROOT / "docs" / "metric_definitions.md").read_text(encoding="utf-8")
    assert k[2023] == 1.0 and k[2024] == 1.05 and k[2025] == 1.07
    assert svi["2024"] == 1.05 and svi["2025"] == 1.1235
    assert "× 1.05" in text and "× 1.1235" in text


def test_price_increase_uniform_across_continuing_lines():
    for r in table("4a_uplift_factors"):
        assert float(r["share_of_lines_at_k_pct"]) == 100.0, r["year"]


def test_margin_allocation_adds_back_to_total_cost(con):
    by_year: dict[int, float] = {}
    for r in table("4f_margin_by_size"):
        by_year[int(r["year"])] = by_year.get(int(r["year"]), 0.0) + float(r["cost_basis_a"])
    for y, cost in by_year.items():
        assert close(cost, scalar(con, f"SELECT sum(total_cost_of_delivery) FROM clean_costs WHERE year(month) = {y}")), y


def test_mix_plus_rate_equals_margin_change():
    for r in table("4f_mix_rate"):
        assert close(float(r["mix_effect_pts"]) + float(r["rate_effect_pts"]), r["change_pts"], 1e-9)


def test_key_figures_trace_to_the_database(con):
    kf = {r["name"]: float(r["value"]) for r in table("key_figures")}
    for y in (2022, 2023, 2024, 2025):
        assert close(kf[f"net_revenue_{y}"], scalar(con, f"SELECT sum(net_revenue) FROM fact_revenue_monthly WHERE year(month_start) = {y}"))
        assert close(kf[f"arr_{y}"], scalar(con, f"SELECT 12 * sum(mrr) FROM fact_mrr_monthly WHERE month_start = DATE '{y}-12-01'"))
        assert close(kf[f"mgmt_revenue_{y}"], scalar(con, f"SELECT sum(total_revenue) FROM clean_management_accounts WHERE year(month) = {y}"))
    assert close(kf["other_implementation_2025"],
                 scalar(con, "SELECT sum(amount) FROM clean_invoices WHERE oneoff_category = 'other_implementation'"))
    for r in table("key_figures"):
        src = r["source"]
        if src.startswith("derived:"):
            # "derived: a - b": recompute from the other key figures
            a_, b_ = [x.strip() for x in src.split(":", 1)[1].split(" - ")]
            assert close(float(r["value"]), kf[a_] - kf[b_]), r["name"]
        else:
            assert (TABLES_DIR / f"{src.split(' ')[0]}.csv").exists() or src.startswith(("clean_", "raw_")), r["name"]


def test_step4_charts_exist():
    charts = ROOT / "outputs" / "charts"
    for name in ["4a_revenue_mix", "4b_pareto", "4b_top10_trend", "4c_logo_retention_heatmap", "4c_revenue_retention_heatmap",
                 "4c_tenure_matched", "4d_nrr_grr", "4e_arr_bridge", "4e_revenue_bridge", "4f_margin", "4f_discount_by_size",
                 "4g_segments", "4g_segment_retention"]:
        path = charts / f"{name}.png"
        if not path.exists():
            pytest.skip("charts not built; run the notebooks")
        assert path.stat().st_size > 10_000, name


def test_qa_references_in_docs_point_at_the_right_question(con):
    assert scalar(con, "SELECT topic FROM qa_log WHERE qa_id = 'Q15'") == "Price increases"


# ---------------------------------------------------------------------------
# Step 5: databook
# ---------------------------------------------------------------------------

DATABOOK = ROOT / "databook" / "Northbridge_Databook.xlsx"


@pytest.fixture(scope="session")
def databook():
    if not DATABOOK.exists():
        pytest.skip("databook not built; run `python run_all.py`")
    from openpyxl import load_workbook
    return load_workbook(DATABOOK, data_only=True), load_workbook(DATABOOK, data_only=False)


def test_databook_recalculated_without_errors():
    report = json.loads((ROOT / "databook" / "recalc_report.json").read_text())
    assert report.get("status") == "success" and report["total_errors"] == 0, report
    assert report["total_formulas"] > 500


def test_databook_every_check_true(databook):
    values, formulas = databook
    ws = values["Checks"]
    assert ws["C5"].value is True
    results = [r[4].value for r in ws.iter_rows(min_row=8) if r[0].value is not None]
    assert len(results) >= 40 and all(v is True for v in results)
    # every check row points at a formula cell that is itself TRUE
    for r in formulas["Checks"].iter_rows(min_row=8):
        if r[0].value is None:
            continue
        sheet, ref = r[2].value, r[3].value
        assert str(formulas[sheet][ref].value).startswith("="), (sheet, ref)
        assert values[sheet][ref].value is True, (sheet, ref)


def test_databook_key_figures_match_outputs_to_the_penny(databook):
    values, _ = databook
    cell_map = json.loads((ROOT / "databook" / "cell_map.json").read_text())
    checked = 0
    for key, m in cell_map.items():
        if key == "overall_check":
            continue
        rows = table(m["table"])
        for col_, val in m["where"].items():
            rows = [r for r in rows if str(r[col_]) == str(val) or (r[col_].replace(".0", "") == str(val))]
        assert len(rows) == 1, key
        expected = float(rows[0][m["column"]]) * m["scale"]
        got = values[m["sheet"]][m["cell"]].value
        assert got is not None and abs(float(got) - expected) <= m["tol"], (key, got, expected)
        checked += 1
    assert checked >= 100


def test_databook_inputs_blue_and_formulas_black(databook):
    _, formulas = databook
    blue = ("FF0000FF", "000000FF", "0000FF")
    for name in ["Reconciliation", "4a Revenue quality", "4b Concentration", "4d NRR GRR", "4e ARR bridge", "4f Margin", "4g Segments"]:
        ws = formulas[name]
        n_inputs = n_formulas = 0
        for row in ws.iter_rows(min_row=5):
            for c in row:
                colour = c.font.color.rgb if c.font and c.font.color and isinstance(c.font.color.rgb, str) else None
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert colour not in blue, (name, c.coordinate)
                    n_formulas += 1
                elif isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    assert colour in blue, (name, c.coordinate, c.value)
                    n_inputs += 1
        assert n_inputs > 0 and n_formulas > 0, name


def test_databook_has_native_charts_and_draft_commentary(databook):
    _, formulas = databook
    for name in ["Reconciliation", "4a Revenue quality", "4b Concentration", "4c Cohorts", "4d NRR GRR", "4e ARR bridge", "4f Margin", "4g Segments"]:
        ws = formulas[name]
        assert len(ws._charts) >= 1, name
        assert any(isinstance(c.value, str) and c.value.startswith("DRAFT commentary") for row in ws.iter_rows() for c in row), name


# ---------------------------------------------------------------------------
# Step 6: Power BI expected values
# ---------------------------------------------------------------------------

def expected_values() -> list[dict]:
    path = ROOT / "dashboard" / "expected_values.csv"
    if not path.exists():
        pytest.skip("expected_values.csv not built; run `python run_all.py`")
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def test_expected_values_agree_with_step4():
    ev = {(int(r["year"]), r["measure"]): r["value"] for r in expected_values() if r["context"] == "year"}
    mix = {int(r["year"]): r for r in table("4a_revenue_mix")}
    arr = {int(r["year"]): r for r in table("4a_arr")}
    nrr = {int(r["year"]): r for r in table("4d_nrr_grr") if r["basis"] == "reported"}
    top = {int(r["year"]): r for r in table("4b_top_n_shares") if r["top_n"] == "10"}
    gm = {int(r["year"]): r for r in table("4f_margin_reconciliation")}
    kf = {r["name"]: float(r["value"]) for r in table("key_figures")}
    for y in (2022, 2023, 2024, 2025):
        assert close(ev[(y, "Net Revenue")], mix[y]["net_revenue"])
        assert close(ev[(y, "Recurring Revenue")], mix[y]["recurring"])
        assert close(ev[(y, "One-off Revenue")], mix[y]["oneoff"])
        assert close(ev[(y, "ARR")], arr[y]["arr"])
        assert int(float(ev[(y, "Active Customers")])) == int(float(arr[y]["active_customers"]))
        assert close(ev[(y, "Top 10 Concentration %")], float(top[y]["share_of_net_pct"]) / 100, 1e-6)
        assert close(ev[(y, "Gross Margin %")], float(gm[y]["gross_margin_pct"]) / 100, 1e-6)
        if y >= 2023:
            assert close(ev[(y, "NRR")], float(nrr[y]["nrr_pct"]) / 100, 1e-6)
            assert close(ev[(y, "GRR")], float(nrr[y]["grr_pct"]) / 100, 1e-6)
            assert close(ev[(y, "Logo Churn Rate")], kf[f"logo_churn_rate_{y}"] / 100, 1e-6)
        else:
            assert ev[(y, "NRR")] == "" and ev[(y, "GRR")] == ""   # no December 2021


def test_expected_values_segments_agree_with_4g():
    seg = {(r["dimension"], r["segment"]): r for r in table("4g_segments")}
    for r in expected_values():
        if r["context"] in ("2025 x region_group", "2025 x size_band") and r["measure"] in ("NRR", "Logo Churn Rate", "Net Revenue"):
            dim = "region_group" if r["filter_column"] == "region_group" else "company_size"
            s = seg[(dim, r["filter_value"])]
            col_, scale = {"NRR": ("nrr_2025_pct", 0.01), "Logo Churn Rate": ("logo_churn_2025_pct", 0.01),
                           "Net Revenue": ("net_revenue_2025", 1)}[r["measure"]]
            assert close(r["value"], float(s[col_]) * scale, 1e-6 if scale != 1 else 0.01), (r["filter_value"], r["measure"])


def test_expected_values_additive_measures_add_up():
    rows = expected_values()
    total = {(int(r["year"]), r["measure"]): float(r["value"]) for r in rows if r["context"] == "year" and r["value"]}
    for ctx, years in (("year x product line", (2022, 2023, 2024, 2025)), ("2025 x region_group", (2025,)),
                       ("2025 x region", (2025,)), ("2025 x size_band", (2025,))):
        for y in years:
            for m in ("Net Revenue", "Recurring Revenue", "ARR"):
                parts = sum(float(r["value"]) for r in rows if r["context"] == ctx and int(r["year"]) == y and r["measure"] == m)
                assert close(parts, total[(y, m)]), (ctx, y, m)


def test_model_data_exported():
    model = ROOT / "dashboard" / "model_data"
    for name in ["FactRevenue", "FactMRR", "FactCost", "DimCustomer", "DimProduct", "DimProductLine", "DimDate"]:
        assert (model / f"{name}.csv").exists(), name


# ---------------------------------------------------------------------------
# Step 7: every £ and % in the written outputs comes from outputs/tables/
# ---------------------------------------------------------------------------

import re
from bisect import bisect_left


def pipeline_values() -> list[float]:
    vals: set[float] = set()
    for path in TABLES_DIR.glob("*.csv"):
        with open(path, newline="", encoding="utf-8") as f:
            for row in csv.reader(f):
                for cell in row:
                    try:
                        v = float(cell)
                    except ValueError:
                        continue
                    vals.update({v, abs(v), 100 * v, abs(100 * v)})
    return sorted(vals)


def has_value(values: list[float], target: float, tol: float) -> bool:
    i = bisect_left(values, target - tol)
    return i < len(values) and values[i] <= target + tol


def figures_in(text: str) -> list[tuple[str, float, float]]:
    """(token, value, tolerance) for every £ and % figure in a text."""
    out = []
    for mo in re.finditer(r"£(\d[\d,]*(?:\.\d+)?)m\b", text):
        s = mo.group(1).replace(",", "")
        dp = len(s.split(".")[1]) if "." in s else 0
        out.append((mo.group(0), float(s) * 1e6, 0.5 * 10 ** -dp * 1e6 + 1e-6))
    for mo in re.finditer(r"£(\d[\d,]*(?:\.\d+)?)(?![\d,.]*m\b)", text):
        s = mo.group(1).replace(",", "")
        dp = len(s.split(".")[1]) if "." in s else 0
        out.append((mo.group(0), float(s), 0.5 * 10 ** -dp + 1e-9))
    for mo in re.finditer(r"(?<![\w.])(\d+(?:\.\d+)?)%", text):
        s = mo.group(1)
        dp = len(s.split(".")[1]) if "." in s else 0
        out.append((mo.group(0), float(s), 0.5 * 10 ** -dp + 1e-9))
    return out


@pytest.mark.parametrize("doc", ["memo/findings_memo.md", "outputs/key_findings.md"])
def test_written_figures_match_outputs(doc):
    path = ROOT / doc
    if not path.exists():
        pytest.skip(f"{doc} not built")
    values = pipeline_values()
    figures = figures_in(path.read_text(encoding="utf-8"))
    assert len(figures) >= 30, doc
    missing = [tok for tok, v, tol in figures if not has_value(values, v, tol)]
    assert not missing, f"figures in {doc} not found in outputs/tables: {missing}"


def test_memo_is_two_pages_or_fewer():
    pdf = ROOT / "memo" / "findings_memo.pdf"
    if not pdf.exists():
        pytest.skip("memo PDF not built")
    from pypdf import PdfReader
    assert len(PdfReader(str(pdf)).pages) <= 2


def test_memo_has_required_sections():
    text = (ROOT / "memo" / "findings_memo.md").read_text(encoding="utf-8")
    for heading in ("1. Scope and basis", "2. Key findings", "3. Red flags", "4. How the investor should view revenue",
                    "5. Questions for management", "6. Overall view"):
        assert heading in text, heading
    questions = re.findall(r"^\d+\. \(Q\d{2}\)", text, flags=re.M)
    assert 10 <= len(questions) <= 15
    assert "synthetic" in text.lower()
    for mo in re.finditer(r"\((Q\d{2})\)", text):   # every question id cited exists in the Q&A log
        assert mo.group(1) in {r["qa_id"] for r in table("qa_log")}
